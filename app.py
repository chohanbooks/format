import io
import copy
import streamlit as st
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

st.set_page_config(page_title="Word Format Copier", page_icon="📄", layout="centered")

st.title("📄 Word Format Copier")
st.write(
    "Upload one correctly formatted Word document as a reference and a second "
    "document whose formatting you want to improve."
)
st.info(
    "This first version copies common styles and page settings. Complex layouts, "
    "section breaks, text boxes, tracked changes, and unusual table structures "
    "may need manual review."
)

reference_file = st.file_uploader(
    "1. Upload the correctly formatted reference file (.docx)",
    type=["docx"],
    key="reference",
)
target_file = st.file_uploader(
    "2. Upload the document to format (.docx)",
    type=["docx"],
    key="target",
)

preserve_content = st.checkbox("Preserve the target document's text and content", value=True)
format_tables = st.checkbox("Apply formatting to tables too", value=True)

def copy_run_format(source_run, target_run):
    """Copy character formatting without changing the target run's text."""
    try:
        if source_run.style is not None:
            target_run.style = source_run.style
    except Exception:
        pass
    try:
        if source_run.font.name:
            target_run.font.name = source_run.font.name
        if source_run.font.size:
            target_run.font.size = source_run.font.size
        target_run.bold = source_run.bold
        target_run.italic = source_run.italic
        target_run.underline = source_run.underline
        if source_run.font.color and source_run.font.color.rgb:
            target_run.font.color.rgb = source_run.font.color.rgb
        if source_run.font.highlight_color:
            target_run.font.highlight_color = source_run.font.highlight_color
        if source_run.font.all_caps is not None:
            target_run.font.all_caps = source_run.font.all_caps
    except Exception:
        pass

def copy_paragraph_format(source_para, target_para):
    """Copy paragraph-level formatting while preserving target paragraph text."""
    try:
        target_para.style = source_para.style
    except Exception:
        pass

    sf = source_para.paragraph_format
    tf = target_para.paragraph_format
    attrs = [
        "alignment", "left_indent", "right_indent", "first_line_indent",
        "space_before", "space_after", "line_spacing", "keep_together",
        "keep_with_next", "page_break_before", "widow_control"
    ]
    for attr in attrs:
        try:
            value = getattr(sf, attr)
            if value is not None:
                setattr(tf, attr, value)
        except Exception:
            pass

    # Apply character formatting to corresponding runs when counts match.
    # If run counts differ, use the first reference run as a sensible fallback.
    if source_para.runs and target_para.runs:
        if len(source_para.runs) == len(target_para.runs):
            for sr, tr in zip(source_para.runs, target_para.runs):
                copy_run_format(sr, tr)
        else:
            source_run = source_para.runs[0]
            for tr in target_para.runs:
                copy_run_format(source_run, tr)

def copy_page_settings(reference_doc, target_doc):
    """Copy common page settings for corresponding sections."""
    for i, target_section in enumerate(target_doc.sections):
        source_section = reference_doc.sections[min(i, len(reference_doc.sections) - 1)]
        for attr in [
            "page_width", "page_height", "top_margin", "bottom_margin",
            "left_margin", "right_margin", "gutter", "header_distance",
            "footer_distance"
        ]:
            try:
                value = getattr(source_section, attr)
                if value is not None:
                    setattr(target_section, attr, value)
            except Exception:
                pass
        try:
            target_section.orientation = source_section.orientation
        except Exception:
            pass

def copy_style_definitions(reference_doc, target_doc):
    """Copy common named style properties into same-named target styles."""
    copied = 0
    for source_style in reference_doc.styles:
        try:
            target_style = target_doc.styles[source_style.name]
        except Exception:
            continue
        try:
            target_style.font.name = source_style.font.name
            target_style.font.size = source_style.font.size
            target_style.font.bold = source_style.font.bold
            target_style.font.italic = source_style.font.italic
            if source_style.font.color and source_style.font.color.rgb:
                target_style.font.color.rgb = source_style.font.color.rgb
        except Exception:
            pass
        try:
            sp, tp = source_style.paragraph_format, target_style.paragraph_format
            for attr in [
                "alignment", "left_indent", "right_indent", "first_line_indent",
                "space_before", "space_after", "line_spacing", "keep_with_next"
            ]:
                value = getattr(sp, attr)
                if value is not None:
                    setattr(tp, attr, value)
        except Exception:
            pass
        copied += 1
    return copied

def format_body_paragraphs(reference_doc, target_doc):
    """Match paragraphs by position, falling back to the last reference paragraph."""
    source_paragraphs = reference_doc.paragraphs
    target_paragraphs = target_doc.paragraphs
    if not source_paragraphs:
        return 0

    count = 0
    for index, target_para in enumerate(target_paragraphs):
        source_para = source_paragraphs[min(index, len(source_paragraphs) - 1)]
        copy_paragraph_format(source_para, target_para)
        count += 1
    return count

def format_tables_by_position(reference_doc, target_doc):
    """Match table formatting by table/cell/paragraph position where possible."""
    count = 0
    if not reference_doc.tables:
        return count

    for ti, target_table in enumerate(target_doc.tables):
        source_table = reference_doc.tables[min(ti, len(reference_doc.tables) - 1)]
        try:
            target_table.alignment = source_table.alignment
            target_table.autofit = source_table.autofit
        except Exception:
            pass

        for ri, target_row in enumerate(target_table.rows):
            if not source_table.rows:
                continue
            source_row = source_table.rows[min(ri, len(source_table.rows) - 1)]
            for ci, target_cell in enumerate(target_row.cells):
                if not source_row.cells:
                    continue
                source_cell = source_row.cells[min(ci, len(source_row.cells) - 1)]
                for pi, target_para in enumerate(target_cell.paragraphs):
                    if source_cell.paragraphs:
                        source_para = source_cell.paragraphs[
                            min(pi, len(source_cell.paragraphs) - 1)
                        ]
                        copy_paragraph_format(source_para, target_para)
                        count += 1
    return count

if st.button("✨ Format Word Document", type="primary", use_container_width=True):
    if not reference_file or not target_file:
        st.error("Please upload both Word documents first.")
    else:
        try:
            reference_doc = Document(io.BytesIO(reference_file.getvalue()))
            target_doc = Document(io.BytesIO(target_file.getvalue()))

            # Preserve the target's text and content. The option is displayed
            # explicitly so users understand this initial version's behavior.
            if not preserve_content:
                st.warning(
                    "This app is designed to preserve target content. "
                    "The current version will still preserve it to avoid accidental data loss."
                )

            copy_page_settings(reference_doc, target_doc)
            style_count = copy_style_definitions(reference_doc, target_doc)
            paragraph_count = format_body_paragraphs(reference_doc, target_doc)

            table_paragraph_count = 0
            if format_tables:
                table_paragraph_count = format_tables_by_position(reference_doc, target_doc)

            output = io.BytesIO()
            target_doc.save(output)
            output.seek(0)

            st.success("Formatting pass completed. Please review the result before using it.")
            st.write(f"- Named styles examined: {style_count}")
            st.write(f"- Body paragraphs processed: {paragraph_count}")
            if format_tables:
                st.write(f"- Table paragraphs processed: {table_paragraph_count}")

            st.download_button(
                label="⬇️ Download formatted Word file",
                data=output.getvalue(),
                file_name="formatted_document.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )
            st.caption(
                "Note: this version matches paragraphs mainly by their order, not by meaning. "
                "For best results, the reference and target documents should have similar structures."
            )
        except Exception as exc:
            st.error("The files could not be processed. Ensure both are valid .docx files.")
            st.exception(exc)

with st.expander("What this version can and cannot do"):
    st.markdown(
        """
        **Included**
        - Common named styles (Normal, Heading 1, Heading 2, etc.)
        - Font family, size, bold, italic, underline, and color where available
        - Paragraph alignment, indentation, spacing, and line spacing
        - Common page dimensions and margins
        - Basic table paragraph formatting

        **Limitations**
        - It matches paragraph formatting mainly by paragraph order. If the two documents have very different structures, results may be imperfect.
        - It does not guarantee exact replication of headers/footers, page-number fields, floating images, text boxes, complex tables, section-specific layouts, or all Word XML properties.
        - Always inspect the downloaded document in Microsoft Word or LibreOffice.
        """
    )
